from __future__ import annotations

import csv
from io import BytesIO
import re
from pathlib import Path

from openpyxl import load_workbook

WHO_DRIVER_GENES = frozenset(
    {
        "ASXL1", "BCOR", "BCORL1", "BRAF", "BRCC3", "CALR", "CBL", "CEBPA",
        "CREBBP", "CSF1R", "CSF3R", "CTCF", "CUX1", "DNMT3A", "ETV6", "EZH2",
        "GATA2", "GNAS", "GNB1", "IDH1", "IDH2", "JAK2", "JAK3", "KDM6A",
        "KIT", "KMT2A", "KRAS", "MPL", "MYD88", "NOTCH1", "NRAS", "PHF6",
        "PIGA", "PPM1D", "PRPF40B", "PTEN", "PTPN11", "RAD21", "RUNX1",
        "SETBP1", "SF1", "SF3A1", "SF3B1", "SMC1A", "SMC3", "SRSF2", "STAG2",
        "STAT3", "TET2", "TP53", "U2AF1", "U2AF2", "WT1", "ZRSR2",
    }
)

_GENE_HEADERS = frozenset({"gene", "gen", "symbol", "driver gen", "drivergen"})


def load_who_driver_genes(path: str | Path | None = None) -> frozenset[str]:
    """Load a local WHO gene list, or use the bundled list when unset."""
    if not path or not str(path).strip():
        return WHO_DRIVER_GENES
    source = Path(path).expanduser()
    if not source.is_file():
        raise ValueError(f"WHO-drivergenfilen finnes ikke: {source}")
    suffix = source.suffix.casefold()
    if suffix == ".xlsx":
        try:
            # Read formulas explicitly: accepting only cached values can
            # silently drop uncalculated reference entries.
            workbook = load_workbook(BytesIO(source.read_bytes()), read_only=True, data_only=False)
            try:
                sheet = workbook["WHO"] if "WHO" in workbook.sheetnames else workbook.active
                records = list(enumerate(sheet.iter_rows(values_only=True), start=1))
            finally:
                workbook.close()
        except Exception as exc:
            raise ValueError(
                f"WHO-drivergenfilen kunne ikke leses som Excel: {source} "
                f"({type(exc).__name__})."
            ) from exc
    elif suffix in {".csv", ".txt"}:
        with source.open(encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(4096)
            stream.seek(0)
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
            except csv.Error:
                delimiter = ","  # Headerless, one-symbol-per-line lists.
            reader = csv.reader(stream, delimiter=delimiter, strict=True)
            records = []
            row_number = 1
            try:
                for row in reader:
                    records.append((row_number, row))
                    row_number = reader.line_num + 1
            except csv.Error as exc:
                raise ValueError(f"WHO-drivergenfilen: Rad {row_number} har ugyldig CSV-format.") from exc
    else:
        raise ValueError("WHO-drivergenfilen må være .xlsx, .csv eller .txt.")

    records = [
        (number, row) for number, row in records
        if any(value is not None and str(value).strip() for value in row)
    ]
    if not records:
        raise ValueError(f"WHO-drivergenfilen er tom: {source}")
    header = [str(value or "").strip().casefold() for value in records[0][1]]
    column = next((index for index, value in enumerate(header) if value in _GENE_HEADERS), 0)
    if header[column] in _GENE_HEADERS:
        records = records[1:]
    parsed_genes = set()
    for row_number, row in records:
        value = row[column] if len(row) > column else None
        symbol = str(value or "").strip().upper()
        if symbol.startswith("="):
            raise ValueError(f"WHO-drivergenfilen: Rad {row_number} inneholder en formel; bruk et bokstavelig gensymbol.")
        if not isinstance(value, str) or not re.fullmatch(r"[A-Z][A-Z0-9-]*", symbol):
            raise ValueError(f"WHO-drivergenfilen: Rad {row_number} har et tomt eller ugyldig gensymbol: {value!r}.")
        parsed_genes.add(symbol)
    genes = frozenset(parsed_genes)
    if not genes:
        raise ValueError(f"WHO-drivergenfilen inneholder ingen gyldige gensymboler: {source}")
    return genes

