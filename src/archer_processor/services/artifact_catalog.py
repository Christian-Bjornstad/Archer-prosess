from __future__ import annotations

from pathlib import Path
from io import BytesIO
import re
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook

from archer_processor.core.rules import (
    default_artifact_rules,
    normalize_artifact_hgvsc,
    parse_artifact_af_threshold,
)


_HEADERS = {
    "gene": {"gene", "gen", "symbol"},
    "hgvsc": {"hgvsc", "hgvscdna"},
    "max_af": {"maxaf", "artifactthroughaf", "afgrense", "artefaktgjennomaf"},
    "reason": {"reason", "begrunnelse", "årsak", "kommentar"},
}


def load_artifact_rules(
    path: str | Path | None = None, *, fallback: list[dict[str, str]] | None = None,
) -> list[dict[str, str]]:
    """Read an explicitly selected Excel catalog; never hide errors with defaults."""
    if not path or not str(path).strip():
        return [dict(entry) for entry in (default_artifact_rules() if fallback is None else fallback)]
    source = Path(path).expanduser()
    if not source.is_file():
        raise ValueError(f"Artefaktlisten finnes ikke: {source}")
    if source.suffix.casefold() != ".xlsx":
        raise ValueError("Artefaktlisten må være en .xlsx-fil.")
    try:
        book = load_workbook(BytesIO(source.read_bytes()), read_only=True, data_only=False)
    except Exception as exc:
        # openpyxl can use either stdlib XML or lxml; normalize both parser
        # failures at this file-input boundary before they reach a Qt slot.
        raise ValueError(f"Artefaktlisten kan ikke leses: {source.name}. {exc}") from exc
    try:
        sheet = next((book[name] for name in ("Artifacts", "Artefakter") if name in book.sheetnames), book.active)
        rows = sheet.iter_rows()
        header = [re.sub(r"[\s_.-]+", "", str(cell.value or "").strip().casefold()) for cell in next(rows, ())]
        columns = {}
        for field, aliases in _HEADERS.items():
            matches = [index for index, value in enumerate(header) if value in aliases]
            if len(matches) > 1:
                raise ValueError(f"Artefaktlisten har flere kolonner for {field}.")
            if matches:
                columns[field] = matches[0]
        for field, label in (("gene", "Gene"), ("hgvsc", "HGVSc")):
            if field not in columns:
                raise ValueError(f"Artefaktlisten mangler kolonnen {label} på rad 1.")
        entries = []
        seen = set()
        for row_number, row in enumerate(rows, start=2):
            cells = {field: row[index] for field, index in columns.items() if index < len(row)}
            if not any(cell.value is not None and str(cell.value).strip() for cell in cells.values()):
                continue
            prefix = f"Rad {row_number} i {source.name}: "
            if any(cell.data_type in {"f", "e"} for cell in cells.values()):
                raise ValueError(prefix + "formel eller Excel-feil må erstattes med en fast verdi.")
            values = {field: "" if cell.value is None else str(cell.value).strip() for field, cell in cells.items()}
            gene = values["gene"].upper()
            hgvsc = values["hgvsc"]
            if not re.fullmatch(r"[A-Z][A-Z0-9-]*", gene):
                raise ValueError(prefix + "Gene må inneholde ett gensymbol.")
            if not re.fullmatch(r"N[MR]_\d+(?:\.\d+)?:c\.[*\d-][^\s:]+", hgvsc):
                raise ValueError(prefix + "HGVSc må inneholde transkript og cDNA-endring, f.eks. NM_000546.6:c.524G>A.")
            key = (gene, normalize_artifact_hgvsc(hgvsc))
            if key in seen:
                raise ValueError(prefix + "duplikat av samme Gene/HGVSc (transkriptversjon ignoreres).")
            seen.add(key)
            entry = {"gene": gene, "hgvsc": hgvsc, "reason": values.get("reason", "")}
            try:
                limit = parse_artifact_af_threshold(values.get("max_af"))
                af_cell = cells.get("max_af")
                if af_cell and isinstance(af_cell.value, (int, float)) and "%" in af_cell.number_format:
                    if not 0 <= af_cell.value <= 1:
                        raise ValueError("AF-prosenten må være mellom 0% og 100%.")
                if limit is not None:
                    entry["max_af"] = f"{limit * 100:.15g}%"
            except ValueError as exc:
                raise ValueError(prefix + f"Ugyldig AF-grense: {exc}") from exc
            entries.append(entry)
        if not entries:
            raise ValueError(f"Artefaktlisten inneholder ingen regler: {source.name}")
        return entries
    except (BadZipFile, ParseError, KeyError) as exc:
        raise ValueError(f"Artefaktlisten kan ikke leses: {source.name}. {exc}") from exc
    finally:
        book.close()
