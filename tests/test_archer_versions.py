import json
from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from archer_processor.core.processor import VariantProcessor
from archer_processor.core.rules import FilterEngine, default_artifact_rules
from archer_processor.io.tsv_reader import ArcherTsvReader
from archer_processor.reports.excel_report import ExcelReportWriter
from archer_processor.services.processed_workbook import ProcessedWorkbookLoader


FIXTURES = Path(__file__).parent / "fixtures"
LAYOUT = json.loads((FIXTURES / "archer_v7_layout.json").read_text(encoding="utf-8"))


def v7_input(tmp_path):
    row = dict.fromkeys(LAYOUT["columns"], "")
    row.update(Sample="DEMO_VPM_S1", Symbol="EZH2", HGVSc="NM_004456.5:c.404G>T",
               AF=0.1, Depth=1200, AO=120, Type="SNP", Source="Vision",
               **{"ClinVar Significance": "Likely_pathogenic",
                  "Genomic Location": "chr7:148508764", "Ref/Alt Allele": "G / T"})
    path = tmp_path / "v7.tsv"
    # Shuffled columns must still give the template order on export.
    pd.DataFrame([row], columns=list(reversed(row))).to_csv(path, sep="\t", index=False)
    return path


def test_versions_and_clinvar_mapping(tmp_path):
    result = VariantProcessor().process(v7_input(tmp_path), "2026-10-05")
    assert result.archer_version == "v7"
    assert result.variants[0].clinical_significance == "Likely_pathogenic"
    assert result.variants[0].raw["ClinVar Significance"] == "Likely_pathogenic"
    assert VariantProcessor().process(FIXTURES / "sample_variants.tsv", "2026-10-05").archer_version == "v6"


def test_v7_export_and_resume_follow_template(tmp_path):
    result = VariantProcessor().process(v7_input(tmp_path), "2026-10-05")
    path = ExcelReportWriter().write(result, tmp_path / "review.xlsx")
    workbook = load_workbook(path)
    for title, offset in [("With Artifacts", 1), ("Artifacts Removed", 0)]:
        ws = workbook[title]
        headers = [c.value for c in ws[1]]
        assert headers[offset:offset + 108] == LAYOUT["columns"]
        assert headers[offset + 108] == "Run date"
        assert ws.freeze_panes == ("J1" if offset else "I1")
        hidden = {c.value for c in ws[1] if ws.column_dimensions[c.column_letter].hidden}
        assert hidden == set(LAYOUT["hidden_columns"])
    ws = workbook["With Artifacts"]
    assert ws.cell(2, 110).value == "2026_10_05"
    state = ProcessedWorkbookLoader().load(path)
    assert state.result.archer_version == "v7"
    assert state.result.run_date == "2026-10-05"
    assert "Run date" not in state.result.variants[0].raw
    assert state.result.variants[0].clinical_significance == "Likely_pathogenic"
    rewritten = ExcelReportWriter().write(state.result, tmp_path / "rewritten.xlsx")
    assert [c.value for c in load_workbook(rewritten)["With Artifacts"][1]] == [c.value for c in ws[1]]


@pytest.mark.parametrize("transcript", ["NM_015338.5", "NM_015338.6", "NM_015338"])
def test_transcript_versions_preserve_asxl1_threshold(transcript):
    from archer_processor.core.highlights import variant_highlight
    base = ArcherTsvReader().read(FIXTURES / "sample_variants.tsv")[0]
    low = replace(base, symbol="ASXL1", hgvsc=f"{transcript}:c.1934dup", af=0.055)
    high = replace(low, af=0.056, matched_rules=[], warnings=[])
    different = replace(low, hgvsc="NM_999999.6:c.1934dup", matched_rules=[], warnings=[])
    FilterEngine().apply([low, high, different])
    assert [v.decision for v in [low, high, different]] == ["excluded", "included", "included"]
    assert variant_highlight(low) == "artifact_light"


@pytest.mark.parametrize("gene,hgvsc", [
    ("EZH2", "NM_004456.5:c.404G>T"),
    ("FBXW7", "NM_033632.3:c.585-7_585-5del"),
    ("PTEN", "NM_000314.8:c.834C>G"),
    ("RUNX1", "NM_001754.5:c.1267C>T"),
])
def test_v7_artifacts_apply_to_both_versions(gene, hgvsc):
    base = ArcherTsvReader().read(FIXTURES / "sample_variants.tsv")[0]
    variant = replace(base, symbol=gene, hgvsc=hgvsc)
    FilterEngine().apply([variant])
    assert variant.decision == "excluded"


def test_combined_catalog_has_no_transcript_version_duplicates():
    import re
    entries = default_artifact_rules()
    keys = {(e["gene"], re.sub(r"\.\d+(?=:)", "", e["hgvsc"])) for e in entries}
    assert len(entries) == len(keys) == 51


def test_v3_settings_merge_preserves_overrides_and_is_idempotent(tmp_path, monkeypatch):
    from archer_processor.services.settings import AppSettings
    from archer_processor.core.rules import legacy_artifact_rules
    path = tmp_path / "config.json"
    monkeypatch.setattr(AppSettings, "config_path", classmethod(lambda cls: path))
    monkeypatch.setattr("archer_processor.services.settings.credentials.get_saved_password", lambda *args: "")
    monkeypatch.setattr("archer_processor.services.settings.credentials.save_password", lambda *args: None)
    rules = legacy_artifact_rules()
    rules[0]["max_af"] = "4%"
    rules.append({"gene": "EZH2", "hgvsc": "NM_004456.5:c.404G>T", "max_af": "2%", "reason": "Local override"})
    path.write_text(json.dumps({"artifact_catalog_version": 3, "artifact_rules": rules}), encoding="utf-8")
    settings = AppSettings.load()
    assert settings.artifact_catalog_version == 4
    assert len(settings.artifact_rules) == 51
    assert settings.artifact_rules[:40] == rules
    settings.save()
    assert AppSettings.load().artifact_rules == settings.artifact_rules


def test_worker_logs_detected_version(tmp_path):
    from archer_processor.gui.app import ProcessingWorker
    from archer_processor.services.settings import AppSettings
    worker = ProcessingWorker(v7_input(tmp_path), tmp_path / "review.xlsx", "2026-10-05", AppSettings(), False)
    messages, errors = [], []
    worker.status.connect(messages.append)
    worker.failed.connect(errors.append)
    worker.run()
    assert errors == []
    assert "Archer version detected: v7 (TSV headers)" in messages
