from pathlib import Path

from openpyxl import Workbook
import pytest

from archer_processor.core import FilterEngine, VariantRecord, default_artifact_rules


def load_rules(*args, **kwargs):
    from archer_processor.services.artifact_catalog import load_artifact_rules
    return load_artifact_rules(*args, **kwargs)


def write_catalog(tmp_path, rows, headers=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Artifacts"
    sheet.append(headers or ["Gene", "HGVSc", "Artifact through AF", "Reason"])
    for row in rows:
        sheet.append(row)
    path = tmp_path / "artifacts.xlsx"
    book.save(path)
    book.close()
    return path


def test_unset_catalog_preserves_built_in_and_custom_fallback_without_aliasing():
    assert load_rules() == default_artifact_rules()
    custom = [{"gene": "TP53", "hgvsc": "NM_000546.6:c.524G>A", "reason": "Synthetic rule"}]
    loaded = load_rules("", fallback=custom)
    loaded[0]["reason"] = "Changed"
    assert custom[0]["reason"] == "Synthetic rule"
    assert load_rules(None, fallback=[]) == []


def test_excel_catalog_reads_gene_column_aliases_and_real_percentage(tmp_path):
    path = write_catalog(tmp_path, [[" asxl1 ", "NM_015338.5:c.1934dup", 0.055, "Synthetic rule"]],
                         ["Gen", "HGVSc", "AF-grense", "Begrunnelse"])
    assert load_rules(path) == [{"gene": "ASXL1", "hgvsc": "NM_015338.5:c.1934dup", "max_af": "5.5%", "reason": "Synthetic rule"}]


def test_excel_changes_are_read_at_each_load_and_blank_af_means_all_af(tmp_path):
    path = write_catalog(tmp_path, [["TP53", "NM_000546.6:c.524G>A", None, "First"]])
    assert load_rules(path)[0]["reason"] == "First"
    write_catalog(tmp_path, [["RUNX1", "NM_001754.4:c.1265A>G", "5,5%", "Second"]])
    assert load_rules(path) == [{"gene": "RUNX1", "hgvsc": "NM_001754.4:c.1265A>G", "max_af": "5.5%", "reason": "Second"}]


@pytest.mark.parametrize("row, message", [
    (["", "NM_000546.6:c.524G>A", None, ""], "Gene"),
    (["TP 53", "NM_000546.6:c.524G>A", None, ""], "Gene"),
    (["TP53", "", None, ""], "HGVSc"),
    (["TP53", "p.Arg175His", None, ""], "HGVSc"),
    (["TP53", "NM_000546.6:c.524 G>A", None, ""], "HGVSc"),
    (["TP53", "NM_000546.6:c.524G>A", "bad", ""], "AF"),
    (["TP53", "NM_000546.6:c.524G>A", "101%", ""], "AF"),
    (["TP53", "NM_000546.6:c.524G>A", "=5.5/100", ""], "formel"),
    (["=\"TP53\"", "NM_000546.6:c.524G>A", None, ""], "formel"),
])
def test_catalog_rejects_invalid_row_instead_of_silently_ignoring_it(tmp_path, row, message):
    path = write_catalog(tmp_path, [row])
    with pytest.raises(ValueError, match=f"[Rr]ad 2.*{message}"):
        load_rules(path)


def test_duplicate_transcript_versions_are_rejected_with_row_number(tmp_path):
    path = write_catalog(tmp_path, [
        ["TP53", "NM_000546.5:c.524G>A", None, ""],
        ["TP53", "NM_000546.6:c.524G>A", "5%", ""],
    ])
    with pytest.raises(ValueError, match="[Rr]ad 3.*duplikat"):
        load_rules(path)


def test_missing_headers_and_empty_catalog_are_rejected(tmp_path):
    path = write_catalog(tmp_path, [["TP53", "NM_000546.6:c.524G>A"]], ["Unknown", "HGVSc"])
    with pytest.raises(ValueError, match="Gene"):
        load_rules(path)
    path = write_catalog(tmp_path, [])
    with pytest.raises(ValueError, match="ingen.*regler"):
        load_rules(path)


def test_blank_rows_and_unrelated_instruction_columns_do_not_become_rules(tmp_path):
    path = write_catalog(tmp_path, [
        [None, None, None, None, "Use the data table"],
        ["TP53", "NM_000546.6:c.524G>A", None, "Synthetic rule"],
    ])
    assert len(load_rules(path)) == 1


def test_corrupt_or_missing_catalog_is_actionable_not_silent_fallback(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_text("not an Excel file", encoding="utf-8")
    with pytest.raises(ValueError, match="kan ikke leses"):
        load_rules(path)
    with pytest.raises(ValueError, match="finnes ikke"):
        load_rules(tmp_path / "missing.xlsx")


def test_malformed_workbook_xml_is_validation_error_and_releases_file(tmp_path):
    from zipfile import ZipFile
    path = tmp_path / "malformed.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<not-closed>")
    with pytest.raises(ValueError, match="kan ikke leses") as error:
        load_rules(path)
    path.rename(tmp_path / "replacement.xlsx")
    assert error.value.__cause__ is not None


def test_percentage_format_cannot_hide_out_of_range_stored_number(tmp_path):
    path = write_catalog(tmp_path, [["TP53", "NM_000546.6:c.524G>A", 5.5, ""]])
    from openpyxl import load_workbook
    book = load_workbook(path)
    book.active["C2"].number_format = "0.0%"
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="[Rr]ad 2.*AF"):
        load_rules(path)


def test_explicit_empty_filter_list_does_not_reenable_default_artifacts():
    variant = VariantRecord(Path("synthetic.tsv"), 2, "DEMO", "ASXL1", "NM_015338.5:c.1934dup", af=0.01)
    FilterEngine([]).apply([variant])
    assert variant.decision == "included"


def test_zero_and_precise_af_limits_are_not_dropped_or_rounded(tmp_path):
    path = write_catalog(tmp_path, [["ASXL1", "NM_015338.5:c.1934dup", 0, "Synthetic"]])
    from archer_processor.core.rules import production_rules
    assert production_rules(load_rules(path))[0].max_af_inclusive == 0.0
    write_catalog(tmp_path, [["ASXL1", "NM_015338.5:c.1934dup", 0.05500001, "Synthetic"]])
    assert production_rules(load_rules(path))[0].max_af_inclusive == pytest.approx(0.05500001, abs=1e-12)
