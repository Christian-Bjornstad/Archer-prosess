from __future__ import annotations

import hashlib
import json
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest
from PIL import Image, ImageDraw

from archer_processor.core.models import DatabaseEvidence, VariantRecord
from archer_processor.services.database_selection import SELECTION_SHEET, SKIP_HEADER
from archer_processor.services.evidence_audit import EvidenceAuditIndex, write_evidence_audit
from archer_processor.services.processed_workbook import ProcessedWorkbookLoader
from archer_processor.services.variant_identity import genomic_identity


ORIGINAL_TIME = datetime(2026, 9, 2, 10, 15, tzinfo=timezone.utc)


class RecordingHistory:
    def __init__(self):
        self.records = []

    def record(self, variant, evidence, *, context="", captured_at=None):
        self.records.append((variant, evidence, context, captured_at))
        return True


def _variant() -> VariantRecord:
    return VariantRecord(
        source_file=Path("synthetic.tsv"), source_row=2, sample="SYNTHETIC",
        symbol="TESTGENE", hgvsc="NM_999999.1:c.10G>A", transcript="NM_999999.1",
        genomic_location="chr1:10010", ref_allele="G", alt_allele="A",
    )


def _audit(tmp_path, *, database="ClinVar", status="found", raw=None):
    variant = _variant()
    root = tmp_path / "review_browser_evidence"
    directory = root / "patient-001" / database.lower()
    directory.mkdir(parents=True)
    screenshot = directory / "verified.png"
    image = Image.new("RGB", (800, 500), "white")
    ImageDraw.Draw(image).rectangle((20, 20, 780, 300), fill="navy")
    image.save(screenshot)
    evidence = DatabaseEvidence(
        database, status, "Synthetic verified result", accession="TEST-123",
        clinical_significance="Synthetic significance",
        raw={
            "assembly_verified": "GRCh37",
            "identity_verification": {
                "accepted": True,
                "requested": asdict(genomic_identity(variant)),
            },
            "screenshots": [{"path": str(screenshot), "label": "Overview"}],
            **(raw or {}),
        },
    )
    path = write_evidence_audit(directory, database, variant, evidence)
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["written_at"] = ORIGINAL_TIME.isoformat()
    path.write_text(json.dumps(payload), encoding="utf-8")
    return root, path, variant, payload


def _load_audit(loader, root, variant, database="ClinVar"):
    return loader._load_audit(
        EvidenceAuditIndex.build(root), root, 1, database, variant,
    )


def test_backfill_records_verified_audit_with_original_time_without_modifying_files(tmp_path):
    root, audit_path, variant, _ = _audit(tmp_path)
    history = RecordingHistory()
    before = audit_path.read_bytes()

    restored = _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant)

    assert len(history.records) == 1
    recorded_variant, recorded_evidence, context, captured_at = history.records[0]
    assert recorded_variant is variant
    assert recorded_evidence is restored
    assert restored.status == "found"
    assert restored.accession == "TEST-123"
    assert restored.raw["identity_verification"]["accepted"] is True
    assert captured_at == ORIGINAL_TIME
    assert context == ""
    assert audit_path.read_bytes() == before


def test_backfill_prefers_capture_time_over_later_audit_write_time(tmp_path):
    captured_at = "2026-08-01T15:20:00+02:00"
    root, _, variant, _ = _audit(tmp_path, raw={"captured_at": captured_at})
    history = RecordingHistory()

    _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant)

    assert history.records[0][3] == datetime.fromisoformat(captured_at)


@pytest.mark.parametrize("timestamp", ["", "yesterday", "2026-09-02T10:15:00", None, 0])
def test_backfill_rejects_missing_invalid_or_naive_original_time(tmp_path, timestamp):
    root, audit_path, variant, payload = _audit(tmp_path)
    payload["written_at"] = timestamp
    audit_path.write_text(json.dumps(payload), encoding="utf-8")
    history = RecordingHistory()

    restored = _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant)

    assert restored.status == "found"
    assert history.records == []


def test_invalid_explicit_capture_time_does_not_become_fresh_at_later_audit_write_time(tmp_path):
    root, _, variant, _ = _audit(tmp_path, raw={"captured_at": "invalid"})
    history = RecordingHistory()

    _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant)

    assert history.records == []


@pytest.mark.parametrize("status", ["historical", "partial_capture", "not_found", "timeout"])
def test_backfill_never_records_old_fallback_or_incomplete_results(tmp_path, status):
    root, _, variant, _ = _audit(tmp_path, status=status)
    history = RecordingHistory()

    _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant)

    assert history.records == []


def test_backfill_runs_after_legacy_safety_migration(tmp_path):
    root, _, variant, _ = _audit(tmp_path, raw={"assembly_verified": ""})
    history = RecordingHistory()

    restored = _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant)

    assert restored.status == "verification_required"
    assert history.records == []


def test_backfill_rejects_found_result_when_screenshot_no_longer_exists(tmp_path):
    root, _, variant, payload = _audit(tmp_path)
    Path(payload["raw"]["screenshots"][0]["path"]).unlink()
    history = RecordingHistory()

    restored = _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant)

    assert restored.status == "partial_capture"
    assert history.records == []


@pytest.mark.parametrize("mutation", ["legacy_schema", "missing_requested_identity", "wrong_allele"])
def test_backfill_requires_precise_current_schema_audit_identity(tmp_path, mutation):
    root, audit_path, variant, payload = _audit(tmp_path)
    if mutation == "legacy_schema":
        payload["schema_version"] = 2
    elif mutation == "missing_requested_identity":
        payload.pop("requested_identity")
        payload["schema_version"] = 2
    else:
        payload["requested_identity"]["alt_allele"] = "T"
    audit_path.write_text(json.dumps(payload), encoding="utf-8")
    history = RecordingHistory()

    _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant)

    assert history.records == []


def test_backfill_retains_mtbp_cancer_context(tmp_path):
    root, _, variant, _ = _audit(tmp_path, database="MTBP", raw={"cancer_type": "Blood"})
    history = RecordingHistory()

    _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant, "MTBP")

    assert history.records[0][2:] == ("Blood", ORIGINAL_TIME)


def test_backfill_preserves_pending_mtbp_report_without_seeding_history(tmp_path):
    root, _, variant, _ = _audit(
        tmp_path, database="MTBP", raw={"analysis_id": "ARCHER-synthetic", "cancer_type": "Blood"},
    )
    history = RecordingHistory()

    restored = _load_audit(ProcessedWorkbookLoader(history_store=history), root, variant, "MTBP")

    assert restored.raw["analysis_id"] == "ARCHER-synthetic"
    assert history.records == []


def _workbook(path, *, value="[found] Text only"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SELECTION_SHEET
    headers = [
        "Sample", "Symbol", "HGVSc", "AF", "Depth", "AO", "Type", "Source",
        "Genomic Location", "Ref/Alt Allele", "Trans", SKIP_HEADER, "ClinVar Evidence",
    ]
    sheet.append(headers)
    sheet.append([
        "SYNTHETIC", "TESTGENE", "NM_999999.1:c.10G>A", 0.1, 100, 10, "SNP", "Archer",
        "chr1:10010", "G/A", "NM_999999.1", "", value,
    ])
    workbook.save(path)
    workbook.close()


def test_plain_workbook_found_cell_never_seeds_history(tmp_path):
    path = tmp_path / "review.xlsx"
    _workbook(path)
    history = RecordingHistory()
    before = hashlib.sha256(path.read_bytes()).digest()

    state = ProcessedWorkbookLoader(history_store=history).load(path)

    assert state.evidence
    assert history.records == []
    assert hashlib.sha256(path.read_bytes()).digest() == before


def test_hidden_workbook_evidence_without_precise_audit_never_seeds_history(tmp_path):
    path = tmp_path / "review.xlsx"
    _workbook(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.create_sheet("_Evidence")
    sheet.append(["Variant", "Database", "Status", "Summary", "Accession", "Significance", "URL", "Raw"])
    sheet.append([
        "SYNTHETIC|NM_999999.1:c.10G>A", "ClinVar", "found", "Synthetic", "TEST-123", "", "",
        json.dumps({"assembly_verified": "GRCh37", "captured_at": ORIGINAL_TIME.isoformat()}),
    ])
    workbook.save(path)
    workbook.close()
    history = RecordingHistory()

    state = ProcessedWorkbookLoader(history_store=history).load(path)

    assert state.evidence
    assert history.records == []


def test_loader_without_injected_store_restores_audit_normally(tmp_path):
    root, _, variant, _ = _audit(tmp_path)

    restored = _load_audit(ProcessedWorkbookLoader(), root, variant)

    assert restored.status == "found"


def test_backfill_failure_is_nonfatal_and_reaches_processing_warnings(tmp_path):
    from archer_processor.services.evidence_history import EvidenceHistoryError

    class FailingHistory:
        def record(self, *args, **kwargs):
            raise EvidenceHistoryError("Synthetic database unavailable")

    path = tmp_path / "review.xlsx"
    _workbook(path)
    _audit(tmp_path)
    before = path.read_bytes()

    state = ProcessedWorkbookLoader(history_store=FailingHistory()).load(path)

    assert state.evidence["SYNTHETIC|NM_999999.1:c.10G>A"][0].status == "found"
    assert len(state.result.warnings) == 1
    assert "histor" in state.result.warnings[0].casefold()
    assert path.read_bytes() == before


def test_audit_import_into_real_history_preserves_original_time_across_reloads(tmp_path):
    from archer_processor.services.evidence_history import EvidenceHistoryStore

    now = datetime(2026, 10, 15, 13, 0, tzinfo=timezone.utc)
    root, audit_path, variant, payload = _audit(tmp_path)
    store = EvidenceHistoryStore(tmp_path / "history", clock=lambda: now)
    loader = ProcessedWorkbookLoader(history_store=store)
    before = audit_path.read_bytes()

    _load_audit(loader, root, variant)
    _load_audit(loader, root, variant)
    restored = store.fallback(
        variant, DatabaseEvidence("ClinVar", "timeout", "Synthetic page timeout"),
        tmp_path / "replay_output",
    )

    assert restored.status == "historical"
    assert restored.raw["historical_result"]["captured_at"] == ORIGINAL_TIME.isoformat()
    assert restored.raw["historical_result"]["used_at"] == now.isoformat()
    assert restored.clinical_significance == payload["clinical_significance"]
    assert Path(restored.raw["screenshots"][0]["path"]).is_file()
    assert audit_path.read_bytes() == before
