from pathlib import Path

from openpyxl import load_workbook

from archer_processor.core import default_artifact_rules
from archer_processor.reports.who_genes import WHO_DRIVER_GENES, load_who_driver_genes
from archer_processor.services.artifact_catalog import load_artifact_rules


ROOT = Path(__file__).resolve().parents[1] / "reference_lists"


def test_delivered_reference_lists_exactly_match_existing_built_in_content():
    assert load_artifact_rules(ROOT / "Artefaktliste.xlsx") == default_artifact_rules()
    assert load_who_driver_genes(ROOT / "WHO-drivergener.xlsx") == WHO_DRIVER_GENES
    for filename in ("Artefaktliste.xlsx", "WHO-drivergener.xlsx"):
        workbook = load_workbook(ROOT / filename)
        try:
            assert len(workbook.sheetnames) == 1
            assert workbook.active.freeze_panes == "A2"
            assert workbook.active.tables
        finally:
            workbook.close()


def test_additional_records_are_loaded_even_beyond_original_excel_table(tmp_path):
    artifact = tmp_path / "artifacts.xlsx"
    book = load_workbook(ROOT / "Artefaktliste.xlsx")
    book.active.append(["TP53", "NM_000546.6:c.524G>A", 0.125, "Synthetic extension"])
    book.save(artifact)
    book.close()
    assert load_artifact_rules(artifact)[-1] == {
        "gene": "TP53", "hgvsc": "NM_000546.6:c.524G>A", "max_af": "12.5%", "reason": "Synthetic extension",
    }
    who = tmp_path / "who.xlsx"
    book = load_workbook(ROOT / "WHO-drivergener.xlsx")
    book.active.append(["DEMO1"])
    book.save(who)
    book.close()
    assert load_who_driver_genes(who) == WHO_DRIVER_GENES | {"DEMO1"}
